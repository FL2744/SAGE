"""Editable Word export of SAGE's rendered content (optional python-docx)."""
from html.parser import HTMLParser
import hashlib
import os
from pathlib import Path
import tempfile


def require_docx():
    try:
        import docx
    except ImportError as error:
        raise RuntimeError('Word export needs python-docx. From the SAGE folder run: '
                           'python3 -m pip install ".[word]"') from error
    return docx


def prepare_entries(entries):
    from copy import deepcopy
    from .citations import without_citations
    result = deepcopy(entries)
    for entry in result:
        for paragraph in entry["paragraphs"]:
            paragraph["text"] = without_citations(paragraph["text"])
    return result


def bookmark_name(value):
    return 'sage_' + hashlib.sha1(value.encode()).hexdigest()[:32]


class WordParser(HTMLParser):
    """Convert the small HTML vocabulary emitted by book.render, never arbitrary HTML."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        docx = require_docx()
        from docx.shared import Inches, Pt, RGBColor
        self.document = docx.Document()
        section = self.document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.8)
        section.left_margin = section.right_margin = Inches(.85)
        for name, size in [('Normal', 12), ('Title', 28), ('Heading 1', 19), ('Heading 2', 13)]:
            style = self.document.styles[name]
            style.font.name = 'Times New Roman'
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
        from docx.oxml.ns import qn
        for style in self.document.styles:
            if hasattr(style, 'paragraph_format'):
                formatting = style.paragraph_format
                formatting.line_spacing = 1.0
                formatting.space_before = Pt(0)
                formatting.space_after = Pt(0)
                formatting.page_break_before = False
                formatting.keep_with_next = False
                formatting.keep_together = False
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
            for fonts in style.element.iter(qn('w:rFonts')):
                for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                    fonts.attrib.pop(qn('w:' + attr), None)
        normal = self.document.styles['Normal'].paragraph_format
        normal.space_after = Pt(0)
        normal.line_spacing = 1.0
        self.document.core_properties.author = 'SAGE'
        self.document.core_properties.title = ''
        self.stack = []
        self.lists = []
        self.paragraph = None
        self.pending = []
        self.bookmark_id = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        skipped = (any(item[2] for item in self.stack) or tag in ('head', 'footer')
                   or attrs.get('class') in ('eyebrow', 'meta')
                   or (tag == 'p' and any(item[0] == 'header' for item in self.stack))
                   or (tag == 'a' and attrs.get('href') == '#contents'))
        self.stack.append((tag, attrs, skipped))
        if skipped:
            return
        if attrs.get('id'):
            self.pending.append(attrs['id'])
        if tag == 'ol':
            self.lists.append(0)
        if tag in ('h1', 'h2', 'h3', 'p', 'li'):
            style = {'h1': 'Title', 'h2': 'Heading 1', 'h3': 'Heading 2'}.get(tag)
            # Two paragraph breaks: one ends the prior block, one blank line separates it.
            if self.document.paragraphs:
                self.document.add_paragraph()
            self.paragraph = self.document.add_paragraph(style=style)
            if attrs.get('class') == 'bibliography-entry':
                from docx.shared import Inches
                self.paragraph.paragraph_format.left_indent = Inches(.3)
                self.paragraph.paragraph_format.first_line_indent = Inches(-.3)
            for ident in self.pending:
                self.add_bookmark(ident)
            self.pending.clear()
            if tag == 'li' and self.lists:
                self.lists[-1] += 1
                self.paragraph.add_run(str(self.lists[-1]) + '. ')
        elif tag == 'br' and self.paragraph is not None:
            self.paragraph.add_run().add_break()
        if tag in ('meta', 'br', 'hr', 'link', 'img', 'input'):
            self.stack.pop()

    def add_bookmark(self, ident):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        self.bookmark_id += 1
        for kind in ('bookmarkStart', 'bookmarkEnd'):
            element = OxmlElement('w:' + kind)
            element.set(qn('w:id'), str(self.bookmark_id))
            if kind == 'bookmarkStart':
                element.set(qn('w:name'), bookmark_name(ident))
            self.paragraph._p.append(element)

    def handle_endtag(self, tag):
        matches = [i for i, item in enumerate(self.stack) if item[0] == tag]
        if not matches:
            return
        index = matches[-1]
        skipped = self.stack[index][2]
        self.stack = self.stack[:index]
        if not skipped:
            if tag == 'ol':
                self.lists.pop()
            if tag in ('p', 'li', 'h1', 'h2', 'h3'):
                self.paragraph = None

    def handle_data(self, data):
        if self.paragraph is None or any(item[2] for item in self.stack):
            return
        tags = [item[0] for item in self.stack]
        run = self.paragraph.add_run(data)
        run.italic = 'i' in tags or 'em' in tags
        run.bold = 'b' in tags or 'strong' in tags
        run.font.superscript = 'sup' in tags
        links = [item[1].get('href', '') for item in self.stack if item[0] == 'a']
        if links:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            from docx.shared import RGBColor
            url = links[-1]
            hyperlink = OxmlElement('w:hyperlink')
            if url.startswith('#'):
                hyperlink.set(qn('w:anchor'), bookmark_name(url[1:]))
            elif url.startswith(('https://', 'http://')):
                hyperlink.set(qn('r:id'), self.document.part.relate_to(url, RT.HYPERLINK, is_external=True))
            else:
                return
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.underline = True
            hyperlink.append(run._r)
            self.paragraph._p.append(hyperlink)
        if 'h1' in tags:
            self.document.core_properties.title += data


def write(rendered_html, output):
    """Atomically save Word content with headings, a single-column TOC and live links."""
    parser = WordParser()
    parser.feed(rendered_html)
    parser.close()
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(suffix='.docx', dir=output.parent)
    os.close(fd)
    try:
        parser.document.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
